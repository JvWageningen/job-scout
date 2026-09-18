"""Interview preparation as a file the applicant can edit and take along.

Both halves of the interview tab export to Word (.docx, which Word, LibreOffice
and Google Docs all open and edit) and to plain text. The two formats are
rendered from one list of blocks, so they carry the same content in the same
order and cannot drift apart.

The document is written for editing, not for show: every question and every
draft answer is a plain paragraph, never a table cell or a text box, so the
applicant can rewrite it in place. Headings use Word's own heading styles,
which keeps the navigation pane useful and the file tidy in any editor.

The fixed wording is in the set's own language and follows the house style in
:mod:`job_scout.writing_style`: no em or en dashes, no decorative symbols and
no emoji. What the model wrote, and what the applicant rewrote, is exported as
it is; the export never rewrites the applicant's own text. The one exception
is technical: control characters that a Word file cannot hold (they arrive
with text pasted from Word or scraped from a page) are dropped from both
formats alike, so one stray character cannot cost the applicant the download.
"""

from __future__ import annotations

import io
import re
import unicodedata
from datetime import UTC, date, datetime
from enum import StrEnum
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml.ns import qn
from docx.shared import Pt
from pydantic import BaseModel, ConfigDict, field_validator

from job_scout.config import user_db_path
from job_scout.database import Database
from job_scout.interview_answers import (
    AnswerFooting,
    InterviewAnswerSet,
    LikelyQuestion,
    QuestionKind,
)
from job_scout.interview_questions import (
    NO_PUBLIC_INFO,
    NO_REVIEW,
    RESEARCH_FAILED,
    THIN_REVIEW,
    InterviewQuestion,
    QuestionTheme,
)
from job_scout.interview_store import InterviewMode, InterviewSet, mode_of
from job_scout.letters.models import LetterLanguage
from job_scout.letters.writer import require_user


class ExportFormat(StrEnum):
    """The file types a set can be downloaded as."""

    DOCX = "docx"
    TXT = "txt"


MEDIA_TYPES = {
    ExportFormat.DOCX: (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ),
    ExportFormat.TXT: "text/plain; charset=utf-8",
}


class ExportedFile(BaseModel):
    """One rendered download.

    Attributes:
        filename: ASCII-only name, safe to put in a Content-Disposition header.
        media_type: MIME type of ``content``.
        content: The file itself.
    """

    filename: str
    media_type: str
    content: bytes


# Word's manual line break and page break, which paste along with Word text.
_XML_BREAKS = re.compile("[\x0b\x0c]")
# Everything else XML 1.0 forbids: C0 controls other than tab, line feed and
# carriage return, lone surrogates, and the two non-characters.
_XML_ILLEGAL = re.compile("[\x00-\x08\x0e-\x1f\ud800-\udfff\ufffe\uffff]")


class BlockKind(StrEnum):
    """The parts a document is made of, rendered the same way in both formats."""

    TITLE = "title"
    SUBTITLE = "subtitle"
    DATE = "date"
    HEADING = "heading"
    QUESTION = "question"
    DETAIL = "detail"
    ANSWER = "answer"
    ITEM = "item"


def _xml_safe(text: str) -> str:
    """Drop the characters a Word file cannot hold.

    Text pasted from Word or scraped from a page can carry control characters
    that XML forbids, and python-docx refuses the whole document for one of
    them. A vertical tab or form feed (Word's line and page breaks) becomes a
    line break; the rest are removed. Tabs and line breaks stay.

    Args:
        text: Any text headed for the document.

    Returns:
        The text with only characters XML allows.
    """
    return _XML_ILLEGAL.sub("", _XML_BREAKS.sub("\n", text))


class Block(BaseModel):
    """One part of the document.

    Both formats are rendered from blocks, so the characters a Word file
    cannot hold are removed here, once, and the text file stays identical.

    Attributes:
        kind: How the part is laid out.
        text: Its text.
        label: For a detail line or an answer, the caption that introduces it.
    """

    model_config = ConfigDict(frozen=True)

    kind: BlockKind
    text: str
    label: str = ""

    @field_validator("text", "label")
    @classmethod
    def _printable(cls, value: str) -> str:
        """Keep only what a Word file can hold.

        Args:
            value: The text or the label.

        Returns:
            The value without XML-illegal characters.
        """
        return _xml_safe(value)


class _Labels(BaseModel):
    """Every fixed word the export writes, in one language."""

    model_config = ConfigDict(frozen=True)

    at: str
    ask_subtitle: str
    answer_subtitle: str
    generated_on: str
    why: str
    grounded_in: str
    why_asked: str
    answer: str
    footing: str
    based_on: str
    based_on_nothing: str
    missing: str
    sources: str
    ask_file: str
    answer_file: str
    months: tuple[str, ...]
    themes: dict[QuestionTheme, str]
    kinds: dict[QuestionKind, str]
    footings: dict[AnswerFooting, str]
    missing_names: dict[str, str]


_LABELS = {
    LetterLanguage.EN: _Labels(
        at="at",
        ask_subtitle="Questions to ask them",
        answer_subtitle="Questions they may ask you, with draft answers",
        generated_on="Generated on",
        why="Why this matters for you",
        grounded_in="Based on",
        why_asked="Why they may ask this",
        answer="Your answer",
        footing="Footing",
        based_on="Based on",
        based_on_nothing="nothing in your CV, stories or notes",
        missing="What was missing",
        sources="Sources about you that were used",
        ask_file="Interview questions",
        answer_file="Interview answers",
        months=(
            "January",
            "February",
            "March",
            "April",
            "May",
            "June",
            "July",
            "August",
            "September",
            "October",
            "November",
            "December",
        ),
        themes={
            QuestionTheme.ROLE: "The role",
            QuestionTheme.TEAM: "The team",
            QuestionTheme.COMPANY: "The company",
            QuestionTheme.GROWTH: "Growth and future",
            QuestionTheme.WAYS_OF_WORKING: "Ways of working",
            QuestionTheme.CONCERNS: "Worth probing",
        },
        kinds={
            QuestionKind.MOTIVATION: "Motivation and fit",
            QuestionKind.EXPERIENCE: "Your experience",
            QuestionKind.TECHNICAL: "Technical depth",
            QuestionKind.BEHAVIOURAL: "How you work with others",
            QuestionKind.GAP: "Gaps they will probe",
            QuestionKind.PRACTICAL: "Practical matters",
        },
        footings={
            AnswerFooting.STRONG: "Strong. Your CV or a STAR story carries this.",
            AnswerFooting.PARTIAL: (
                "Partial. Only related experience, so choose the framing with care."
            ),
            AnswerFooting.GAP: (
                "Gap. You do not have this, so rehearse saying so plainly."
            ),
        },
        missing_names={},
    ),
    LetterLanguage.NL: _Labels(
        at="bij",
        ask_subtitle="Vragen die jij stelt",
        answer_subtitle="Vragen die zij jou kunnen stellen, met een conceptantwoord",
        generated_on="Gemaakt op",
        why="Waarom dit voor jou telt",
        grounded_in="Gebaseerd op",
        why_asked="Waarom ze dit vragen",
        answer="Jouw antwoord",
        footing="Onderbouwing",
        based_on="Gebaseerd op",
        based_on_nothing="niets uit je cv, verhalen of notities",
        missing="Wat ontbrak",
        sources="Gebruikte bronnen over jou",
        ask_file="Interviewvragen",
        answer_file="Interviewantwoorden",
        months=(
            "januari",
            "februari",
            "maart",
            "april",
            "mei",
            "juni",
            "juli",
            "augustus",
            "september",
            "oktober",
            "november",
            "december",
        ),
        themes={
            QuestionTheme.ROLE: "De functie",
            QuestionTheme.TEAM: "Het team",
            QuestionTheme.COMPANY: "Het bedrijf",
            QuestionTheme.GROWTH: "Groei en toekomst",
            QuestionTheme.WAYS_OF_WORKING: "Manier van werken",
            QuestionTheme.CONCERNS: "Om door te vragen",
        },
        kinds={
            QuestionKind.MOTIVATION: "Motivatie",
            QuestionKind.EXPERIENCE: "Je ervaring",
            QuestionKind.TECHNICAL: "Vakinhoud",
            QuestionKind.BEHAVIOURAL: "Hoe je met anderen werkt",
            QuestionKind.GAP: "Leemtes waar ze naar vragen",
            QuestionKind.PRACTICAL: "Praktische zaken",
        },
        footings={
            AnswerFooting.STRONG: "Sterk. Je cv of een STAR-verhaal onderbouwt dit.",
            AnswerFooting.PARTIAL: (
                "Deels. Alleen verwante ervaring, dus kies je woorden met zorg."
            ),
            AnswerFooting.GAP: (
                "Leemte. Dit heb je niet, oefen hoe je dat eerlijk zegt."
            ),
        },
        missing_names={
            NO_PUBLIC_INFO: "geen openbare informatie over het bedrijf gevonden",
            RESEARCH_FAILED: "het bedrijfsonderzoek kon deze keer niet worden afgerond",
            NO_REVIEW: "nog geen beoordeling van het bedrijf",
            THIN_REVIEW: "de beoordeling van het bedrijf rust op weinig bronnen",
            "no vacancy description": "geen vacaturetekst",
            "no STAR stories saved yet": "nog geen STAR-verhalen opgeslagen",
        },
    ),
}

# The applicant's sources are named in English where they are gathered; a Dutch
# document says the same in Dutch. A name not listed here is left as it is.
_SOURCE_NAMES_NL = (
    (re.compile(r"^CV Builder profile '(.+)'$"), r"CV Builder-profiel '\1'"),
    (re.compile(r"^your own CV \((.+)\)$"), r"je eigen cv (\1)"),
    (re.compile(r"^your extra experience notes$"), "je extra notities over je werk"),
    (
        re.compile(r"^your parsed profile \(including any LinkedIn import\)$"),
        "je uitgelezen profiel (met een eventuele LinkedIn-import)",
    ),
    (re.compile(r"^your profile description$"), "je profielbeschrijving"),
    (re.compile(r"^1 career track\(s\)$"), "1 loopbaanrichting"),
    (re.compile(r"^(\d+) career track\(s\)$"), r"\1 loopbaanrichtingen"),
    (re.compile(r"^1 STAR story$"), "1 STAR-verhaal"),
    (re.compile(r"^(\d+) STAR stories$"), r"\1 STAR-verhalen"),
)

# Dates are the applicant's, not the server's: a set written late in the
# evening in the Netherlands belongs to that day, not to the next UTC one.
_LOCAL_ZONE = "Europe/Amsterdam"
_MAX_NAME = 60
_UNSAFE_NAME = re.compile(r"[^A-Za-z0-9 .,&()+-]+")
_PARAGRAPH_BREAK = re.compile(r"\n\s*\n")
_WORD_LANGUAGE = {LetterLanguage.NL: "nl-NL", LetterLanguage.EN: "en-GB"}


def _local_date(moment: datetime) -> date:
    """Return the calendar date a moment falls on in the Netherlands.

    Args:
        moment: A timestamp; a naive one is taken as already local.

    Returns:
        The local date.
    """
    if moment.tzinfo is None:
        return moment.date()
    try:
        return moment.astimezone(ZoneInfo(_LOCAL_ZONE)).date()
    except ZoneInfoNotFoundError:
        return moment.date()


def _spoken_date(day: date, labels: _Labels) -> str:
    """Write a date the way a person does: 18 september 2026.

    Args:
        day: The date.
        labels: Month names in the set's language.

    Returns:
        The date in words.
    """
    return f"{day.day} {labels.months[day.month - 1]} {day.year}"


def _header(
    interview: InterviewSet, vacancy_title: str | None, labels: _Labels
) -> list[Block]:
    """Build the title, the half this is and the date it was generated.

    Args:
        interview: The set being exported.
        vacancy_title: The vacancy's title, or None when it is gone.
        labels: Fixed words in the set's language.

    Returns:
        The opening blocks.
    """
    title = interview.company
    if vacancy_title:
        title = f"{vacancy_title} {labels.at} {interview.company}"
    ask = mode_of(interview) is InterviewMode.ASK
    day = _spoken_date(_local_date(interview.generated_at), labels)
    return [
        Block(kind=BlockKind.TITLE, text=title),
        Block(
            kind=BlockKind.SUBTITLE,
            text=labels.ask_subtitle if ask else labels.answer_subtitle,
        ),
        Block(kind=BlockKind.DATE, text=f"{labels.generated_on} {day}"),
    ]


def _question_blocks(item: InterviewQuestion, labels: _Labels) -> list[Block]:
    """Lay out one question to ask, with why it matters and its source.

    Args:
        item: The question.
        labels: Fixed words in the set's language.

    Returns:
        The question, then its explanation lines.
    """
    blocks = [Block(kind=BlockKind.QUESTION, text=item.question)]
    if item.why:
        blocks.append(Block(kind=BlockKind.DETAIL, label=labels.why, text=item.why))
    if item.grounded_in:
        label = labels.grounded_in
        blocks.append(Block(kind=BlockKind.DETAIL, label=label, text=item.grounded_in))
    return blocks


def _answer_blocks(item: LikelyQuestion, labels: _Labels) -> list[Block]:
    """Lay out one likely question with its draft answer and its footing.

    Args:
        item: The likely question and its draft answer.
        labels: Fixed words in the set's language.

    Returns:
        The question, why it is asked, the answer, the footing and the sources.
    """
    blocks = [Block(kind=BlockKind.QUESTION, text=item.question)]
    if item.why_asked:
        blocks.append(
            Block(kind=BlockKind.DETAIL, label=labels.why_asked, text=item.why_asked)
        )
    based_on = "; ".join(item.based_on) or labels.based_on_nothing
    return [
        *blocks,
        Block(kind=BlockKind.ANSWER, label=labels.answer, text=item.draft_answer),
        Block(
            kind=BlockKind.DETAIL,
            label=labels.footing,
            text=labels.footings[item.footing],
        ),
        Block(kind=BlockKind.DETAIL, label=labels.based_on, text=based_on),
    ]


def _body(interview: InterviewSet, labels: _Labels) -> list[Block]:
    """Group the set the way the dashboard does, in conversation order.

    Args:
        interview: The set being exported.
        labels: Fixed words in the set's language.

    Returns:
        A heading per group that has entries, each followed by its entries.
    """
    blocks: list[Block] = []
    if isinstance(interview, InterviewAnswerSet):
        for kind in QuestionKind:
            picked = [q for q in interview.questions if q.kind is kind]
            if picked:
                blocks.append(Block(kind=BlockKind.HEADING, text=labels.kinds[kind]))
            for answer in picked:
                blocks.extend(_answer_blocks(answer, labels))
        return blocks
    for theme in QuestionTheme:
        chosen = [q for q in interview.questions if q.theme is theme]
        if chosen:
            blocks.append(Block(kind=BlockKind.HEADING, text=labels.themes[theme]))
        for question in chosen:
            blocks.extend(_question_blocks(question, labels))
    return blocks


def _source_name(source: str, language: LetterLanguage) -> str:
    """Name one of the applicant's sources in the document's language.

    Args:
        source: The name as the generator recorded it.
        language: The set's language.

    Returns:
        The Dutch name when one is known and the set is Dutch, else the name.
    """
    if language is not LetterLanguage.NL:
        return source
    for pattern, replacement in _SOURCE_NAMES_NL:
        if pattern.match(source):
            return pattern.sub(replacement, source)
    return source


def _closing(interview: InterviewSet, labels: _Labels) -> list[Block]:
    """List what the generation lacked and which of the applicant's sources it used.

    Args:
        interview: The set being exported.
        labels: Fixed words in the set's language.

    Returns:
        Up to two short sections; an empty one is left out.
    """
    blocks: list[Block] = []
    if interview.missing_context:
        blocks.append(Block(kind=BlockKind.HEADING, text=labels.missing))
        blocks.extend(
            Block(kind=BlockKind.ITEM, text=labels.missing_names.get(gap, gap))
            for gap in interview.missing_context
        )
    if interview.sources_used:
        blocks.append(Block(kind=BlockKind.HEADING, text=labels.sources))
        blocks.extend(
            Block(kind=BlockKind.ITEM, text=_source_name(source, interview.language))
            for source in interview.sources_used
        )
    return blocks


def interview_blocks(
    interview: InterviewSet, vacancy_title: str | None = None
) -> list[Block]:
    """Lay out a set as the blocks both file formats are rendered from.

    Args:
        interview: A question set or an answer set, as it is on screen.
        vacancy_title: The vacancy's title for the heading; None leaves the
            company on its own.

    Returns:
        The document, top to bottom.
    """
    labels = _LABELS[interview.language]
    return [
        *_header(interview, vacancy_title, labels),
        *_body(interview, labels),
        *_closing(interview, labels),
    ]


def _paragraphs(text: str) -> list[str]:
    """Split an answer into the paragraphs its writer left between blank lines.

    Args:
        text: A draft answer.

    Returns:
        Its non-empty paragraphs, each trimmed.
    """
    return [part.strip() for part in _PARAGRAPH_BREAK.split(text) if part.strip()]


def _text_lines(block: Block, previous: Block | None) -> list[str]:
    """Render one block as plain-text lines, with the blank lines before it.

    Args:
        block: The block to render.
        previous: The block before it, which decides the spacing of a list item.

    Returns:
        The lines, without line endings.
    """
    if block.kind is BlockKind.DETAIL:
        return [f"{block.label}: {block.text}"]
    if block.kind is BlockKind.ANSWER:
        return ["", f"{block.label}:", "\n\n".join(_paragraphs(block.text)), ""]
    if block.kind is BlockKind.HEADING:
        return ["", "", block.text.upper()]
    after_heading = previous is not None and previous.kind is BlockKind.HEADING
    if block.kind is BlockKind.QUESTION or after_heading:
        return ["", block.text]
    return [block.text]


def render_text(blocks: list[Block]) -> str:
    """Render the blocks as plain text any editor can open.

    Headings are in capitals, each question starts after a blank line and each
    draft answer stands on its own between blank lines, so the file reads well
    in Notepad and edits as easily as the Word version.

    Args:
        blocks: The document from :func:`interview_blocks`.

    Returns:
        The text, with Unix line endings and one final newline.
    """
    lines: list[str] = []
    previous: Block | None = None
    for block in blocks:
        lines.extend(_text_lines(block, previous))
        previous = block
    return "\n".join(line.rstrip() for line in lines).strip() + "\n"


def _set_language(document: DocxDocument, language: LetterLanguage) -> None:
    """Mark the text as Dutch or English, so the spelling check matches it.

    Args:
        document: The document being built.
        language: The set's language.
    """
    code = _WORD_LANGUAGE[language]
    for lang in document.styles.element.xpath(
        "w:docDefaults/w:rPrDefault/w:rPr/w:lang"
    ):
        lang.set(qn("w:val"), code)
    document.core_properties.language = code


def _set_properties(
    document: DocxDocument, title: str, language: LetterLanguage, moment: datetime
) -> None:
    """Replace the template's properties with this document's own.

    The python-docx template names itself as author and comment; neither
    belongs in a file the applicant keeps.

    Args:
        document: The document being built.
        title: The document title.
        language: The set's language.
        moment: When the set was generated.
    """
    # python-docx writes the time as UTC without converting it first.
    utc = moment.astimezone(UTC) if moment.tzinfo else moment
    properties = document.core_properties
    properties.title = title
    properties.author = ""
    properties.comments = ""
    properties.last_modified_by = ""
    properties.created = utc
    properties.modified = utc
    properties.revision = 1
    _set_language(document, language)


def _add_detail(document: DocxDocument, block: Block) -> None:
    """Add a captioned line such as the reason for a question.

    Args:
        document: The document being built.
        block: A detail block.
    """
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run(f"{block.label}: ").italic = True
    paragraph.add_run(block.text)


def _add_question(document: DocxDocument, block: Block) -> None:
    """Add a question as its own bold paragraph that stays with its lines.

    Args:
        document: The document being built.
        block: A question block.
    """
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run(block.text).bold = True


def _add_answer(document: DocxDocument, block: Block) -> None:
    """Add a draft answer as plain paragraphs under a short caption.

    The answer is the part the applicant rewrites most, so it is ordinary body
    text: no table, no text box, nothing that gets in the way of editing.

    Args:
        document: The document being built.
        block: An answer block.
    """
    caption = document.add_paragraph()
    caption.paragraph_format.space_before = Pt(6)
    caption.paragraph_format.space_after = Pt(2)
    caption.paragraph_format.keep_with_next = True
    caption.add_run(f"{block.label}:").italic = True
    for text in _paragraphs(block.text) or [""]:
        document.add_paragraph(text).paragraph_format.space_after = Pt(6)


def _add_block(document: DocxDocument, block: Block) -> None:
    """Add one block in the Word style that fits it.

    Args:
        document: The document being built.
        block: The block to add.
    """
    if block.kind is BlockKind.TITLE:
        document.add_heading(block.text, level=0)
    elif block.kind is BlockKind.SUBTITLE:
        document.add_paragraph(block.text, style="Subtitle")
    elif block.kind is BlockKind.HEADING:
        document.add_heading(block.text, level=1)
    elif block.kind is BlockKind.QUESTION:
        _add_question(document, block)
    elif block.kind is BlockKind.DETAIL:
        _add_detail(document, block)
    elif block.kind is BlockKind.ANSWER:
        _add_answer(document, block)
    else:
        document.add_paragraph(block.text).paragraph_format.space_after = Pt(4)


def render_docx(
    blocks: list[Block], language: LetterLanguage, moment: datetime
) -> bytes:
    """Render the blocks as a Word document.

    Args:
        blocks: The document from :func:`interview_blocks`.
        language: The set's language, for the spelling check.
        moment: When the set was generated, for the file's properties.

    Returns:
        The .docx file.
    """
    document = Document()
    title = blocks[0].text if blocks else ""
    _set_properties(document, title, language, moment)
    for block in blocks:
        _add_block(document, block)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _ascii_name(text: str) -> str:
    """Fold a company name to letters a file name and a header can carry.

    Args:
        text: The company name as scraped.

    Returns:
        ASCII letters, digits and a few harmless marks, at most 60 characters.
    """
    folded = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    cleaned = " ".join(_UNSAFE_NAME.sub(" ", folded).split())
    return cleaned[:_MAX_NAME].strip(" .,-")


def export_filename(interview: InterviewSet, file_format: ExportFormat) -> str:
    """Name the download: date, what it is and the company.

    For example "20260918 Interviewvragen Findwhere.docx". The name is ASCII
    only, so it survives every browser, file system and header unchanged.

    Args:
        interview: The set being exported.
        file_format: The file type.

    Returns:
        The file name with its extension.
    """
    labels = _LABELS[interview.language]
    ask = mode_of(interview) is InterviewMode.ASK
    stamp = _local_date(interview.generated_at).strftime("%Y%m%d")
    parts = [
        stamp,
        labels.ask_file if ask else labels.answer_file,
        _ascii_name(interview.company),
    ]
    return " ".join(part for part in parts if part) + f".{file_format.value}"


def vacancy_title(user: str, job_id: int) -> str | None:
    """Look up the vacancy's title for the document heading.

    Args:
        user: Name of an existing user.
        job_id: The vacancy.

    Returns:
        The title, or None when the vacancy is gone or has none.
    """
    job = Database(user_db_path(require_user(user))).get_job(job_id)
    if job is None or not job.title.strip():
        return None
    return job.title.strip()


def export_interview(
    user: str, interview: InterviewSet, file_format: ExportFormat
) -> ExportedFile:
    """Render a set, as it is on screen, as a file to download.

    Args:
        user: Name of an existing user, whose vacancy supplies the title.
        interview: The set, including any answers the applicant rewrote.
        file_format: Word or plain text.

    Returns:
        The file with its name and MIME type.
    """
    blocks = interview_blocks(interview, vacancy_title(user, interview.job_id))
    if file_format is ExportFormat.DOCX:
        content = render_docx(blocks, interview.language, interview.generated_at)
    else:
        content = render_text(blocks).encode("utf-8")
    return ExportedFile(
        filename=export_filename(interview, file_format),
        media_type=MEDIA_TYPES[file_format],
        content=content,
    )
