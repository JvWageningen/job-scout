"""Contracts for saving, reading back and downloading interview sets.

The generators are stubbed, as in the other interview API tests: what is under
test is that a generated set is kept, that the page can read it back without
generating again, and that what is on screen downloads as an editable file.

Every person and employer below is invented. The repository is public, so no
fixture may carry a real name, address or company.
"""

from __future__ import annotations

import io
import re
from collections.abc import Callable
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import pytest
from click.testing import CliRunner
from docx import Document
from fastapi.testclient import TestClient

import job_scout.config as config
from job_scout.cli import cli
from job_scout.database import Database
from job_scout.interview_answers import (
    AnswerFooting,
    InterviewAnswerSet,
    LikelyQuestion,
    QuestionKind,
)
from job_scout.interview_questions import (
    InterviewQuestion,
    InterviewQuestionSet,
    QuestionTheme,
)
from job_scout.interview_store import save_interview_set
from job_scout.letters.models import LetterLanguage
from job_scout.models import JobListing, JobStatus
from job_scout.web.app import create_app
from tests.helpers import FakeLLMClient

USER = "Sam"
OTHER = "Robin"
COMPANY = "Findwhere"
TITLE = "Meetspecialist"
TOKEN = "fictional-test-token"
QUESTION = "Hoe vangt het team de piek rond de audits op?"
DRAFT = "Bij Bureau Kalibra draaide ik de jaarlijkse audit, en dat ging goed."
EDITED = "In mijn eigen woorden: ik plande de audit drie maanden vooruit."
DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
STATIC = Path(__file__).parent.parent / "src/job_scout/web/static"


def question_set(job_id: int, question: str = QUESTION) -> InterviewQuestionSet:
    """Build what a successful question generation returns.

    Args:
        job_id: The vacancy.
        question: The question, so two generations can be told apart.

    Returns:
        A one-question set.
    """
    return InterviewQuestionSet(
        job_id=job_id,
        company=COMPANY,
        language=LetterLanguage.NL,
        questions=[
            InterviewQuestion(
                question=question,
                theme=QuestionTheme.CONCERNS,
                why="De reviews noemen werkdruk.",
                grounded_in="company review: cons",
            )
        ],
        generated_at=datetime(2026, 9, 18, 9, tzinfo=UTC),
        missing_context=["no company review yet"],
    )


def answer_set(job_id: int, answer: str = DRAFT) -> InterviewAnswerSet:
    """Build what a successful answer generation returns.

    Args:
        job_id: The vacancy.
        answer: The draft answer.

    Returns:
        A one-question set with a draft answer.
    """
    return InterviewAnswerSet(
        job_id=job_id,
        company=COMPANY,
        language=LetterLanguage.NL,
        questions=[
            LikelyQuestion(
                question="Hoe pak je een jaarlijkse audit aan?",
                kind=QuestionKind.EXPERIENCE,
                why_asked="De vacature noemt de audits.",
                draft_answer=answer,
                based_on=["STAR story 3"],
                footing=AnswerFooting.STRONG,
            )
        ],
        generated_at=datetime(2026, 9, 18, 9, tzinfo=UTC),
    )


def returning(result: Any) -> Callable[..., Any]:
    """Stand in for a generator that always returns the same set.

    Args:
        result: What every call returns.

    Returns:
        A drop-in replacement for either generator.
    """

    def fake(*args: object, **kwargs: object) -> Any:
        return result

    return fake


@pytest.fixture
def job_id(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> int:
    """Two users, one of whom has a vacancy an interview is prepared for.

    Args:
        tmp_path: Private data root for this test.
        monkeypatch: Used to redirect config and the LLM client factory.

    Returns:
        The id of the vacancy belonging to ``USER``.
    """
    monkeypatch.setattr(config, "DATA_DIR", tmp_path)
    monkeypatch.setattr(config, "CONFIG_PATH", tmp_path / "config.yaml")
    config.write_global_config({"llm_provider": "local"})
    for user in (USER, OTHER):
        config.save_user_config(user, {})
    monkeypatch.setattr(
        "job_scout.web.app.get_llm_client", lambda _: FakeLLMClient(["{}"])
    )
    return Database(config.user_db_path(USER)).save_job(
        JobListing(
            title=TITLE,
            company=COMPANY,
            url="https://voorbeeld.example/vacatures/7",
            source="board",
            description="Je bewaakt de kwaliteit van onze meetmethoden.",
            fit_score=80,
            status=JobStatus.MATCHED,
            seen_at=datetime(2026, 1, 1, tzinfo=UTC),
        )
    )


@pytest.fixture
def client() -> TestClient:
    """A dashboard client without a token, like a local install.

    Returns:
        The test client.
    """
    return TestClient(create_app())


def test_a_generated_question_set_is_saved_and_read_back(
    job_id: int, client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Generating once is enough: the page reads the set back afterwards."""
    monkeypatch.setattr(
        "job_scout.web.app.generate_interview_questions",
        returning(question_set(job_id)),
    )

    response = client.post(
        f"/api/interview/questions?user={USER}", json={"job_id": job_id}
    )
    saved = client.get(f"/api/interview/saved/{job_id}?user={USER}")

    assert response.status_code == 200, response.text
    assert response.headers["x-interview-saved"] == "true"
    assert saved.status_code == 200, saved.text
    assert saved.headers["cache-control"] == "no-store"
    body = saved.json()
    assert body["job_id"] == job_id
    assert [s["questions"][0]["question"] for s in body["questions"]] == [QUESTION]
    assert body["questions"][0]["generated_at"].startswith("2026-09-18T09:00:00")
    assert body["answers"] == []


def test_generating_again_replaces_the_saved_set(
    job_id: int, client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """One saved set per vacancy, half and language: the newest one."""
    for question in ("Eerste vraag?", "Tweede vraag?"):
        monkeypatch.setattr(
            "job_scout.web.app.generate_interview_questions",
            returning(question_set(job_id, question)),
        )
        client.post(f"/api/interview/questions?user={USER}", json={"job_id": job_id})

    body = client.get(f"/api/interview/saved/{job_id}?user={USER}").json()

    assert [s["questions"][0]["question"] for s in body["questions"]] == [
        "Tweede vraag?"
    ]


def test_a_generated_answer_set_is_saved_in_its_own_half(
    job_id: int, client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The two halves are kept apart, and the mode filter reads one of them."""
    monkeypatch.setattr(
        "job_scout.web.app.generate_interview_answers", returning(answer_set(job_id))
    )
    save_interview_set(USER, question_set(job_id))

    client.post(f"/api/interview/answers?user={USER}", json={"job_id": job_id})
    both = client.get(f"/api/interview/saved/{job_id}?user={USER}").json()
    only = client.get(f"/api/interview/saved/{job_id}?user={USER}&mode=answer").json()

    assert both["answers"][0]["questions"][0]["draft_answer"] == DRAFT
    assert len(both["questions"]) == 1
    assert only["questions"] == [] and len(only["answers"]) == 1


def test_a_set_that_cannot_be_saved_still_reaches_the_page(
    job_id: int, client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Minutes of generation are not thrown away because the disk refused."""
    monkeypatch.setattr(
        "job_scout.web.app.generate_interview_questions",
        returning(question_set(job_id)),
    )

    def refuse(*args: object) -> None:
        raise OSError("disk full")

    monkeypatch.setattr("job_scout.web.app.save_interview_set", refuse)

    response = client.post(
        f"/api/interview/questions?user={USER}", json={"job_id": job_id}
    )

    assert response.status_code == 200
    assert response.headers["x-interview-saved"] == "false"
    assert response.json()["questions"][0]["question"] == QUESTION


def test_nothing_saved_reads_as_empty_halves(job_id: int, client: TestClient) -> None:
    """A vacancy without saved sets is not an error; the page offers to generate."""
    body = client.get(f"/api/interview/saved/{job_id}?user={USER}").json()

    assert body == {"job_id": job_id, "questions": [], "answers": []}


def test_one_users_saved_sets_never_reach_another(
    job_id: int, client: TestClient
) -> None:
    """The same vacancy number in another user's data is another vacancy."""
    save_interview_set(USER, question_set(job_id))

    theirs = client.get(f"/api/interview/saved/{job_id}?user={OTHER}").json()

    assert theirs["questions"] == []


@pytest.mark.parametrize(
    ("path", "expected"),
    [
        ("/api/interview/saved/0?user=Sam", 400),
        ("/api/interview/saved/-4?user=Sam", 400),
        ("/api/interview/saved/1?user=../Sam", 400),
        ("/api/interview/saved/1?user=Nobody", 400),
        ("/api/interview/saved/1?user=Sam&mode=../../x", 422),
        ("/api/interview/saved/abc?user=Sam", 422),
    ],
)
def test_a_bad_request_never_builds_a_path(
    job_id: int, client: TestClient, path: str, expected: int
) -> None:
    """Every part of the file name is checked before it touches the disk."""
    assert client.get(path).status_code == expected


def test_edited_answers_are_saved_in_place_of_the_drafts(
    job_id: int, client: TestClient
) -> None:
    """The rewrite is what the page shows next time, and what exports."""
    save_interview_set(USER, answer_set(job_id))
    edited = answer_set(job_id, EDITED).model_dump(mode="json")

    response = client.put(
        f"/api/interview/saved/answers/{job_id}?user={USER}", json=edited
    )
    body = client.get(f"/api/interview/saved/{job_id}?user={USER}").json()

    assert response.status_code == 200, response.text
    assert body["answers"][0]["questions"][0]["draft_answer"] == EDITED


def test_a_file_saved_in_another_encoding_hides_only_itself(
    job_id: int, client: TestClient
) -> None:
    """One unreadable file must not fail the whole vacancy with a 400."""
    save_interview_set(USER, question_set(job_id))
    path = save_interview_set(USER, answer_set(job_id, "In het café oefende ik."))
    path.write_bytes(path.read_text(encoding="utf-8").encode("cp1252"))

    response = client.get(f"/api/interview/saved/{job_id}?user={USER}")

    assert response.status_code == 200, response.text
    assert response.json()["answers"] == []
    assert len(response.json()["questions"]) == 1


def test_an_edit_without_a_time_zone_keeps_the_saved_list_readable(
    job_id: int, client: TestClient
) -> None:
    """A bare timestamp beside a zoned one used to make sorting fail with a 500."""
    english = answer_set(job_id).model_copy(update={"language": LetterLanguage.EN})
    save_interview_set(USER, english)
    bare = answer_set(job_id, EDITED).model_dump(mode="json")
    bare["generated_at"] = "2026-09-18T10:00:00"

    put = client.put(f"/api/interview/saved/answers/{job_id}?user={USER}", json=bare)
    response = client.get(f"/api/interview/saved/{job_id}?user={USER}")

    assert put.status_code == 200, put.text
    assert response.status_code == 200, response.text
    answers = response.json()["answers"]
    assert [s["language"] for s in answers] == ["nl", "en"]
    assert answers[0]["generated_at"] == "2026-09-18T10:00:00Z"


def test_a_vacancy_that_left_the_shortlist_keeps_its_saved_sets(
    job_id: int, client: TestClient
) -> None:
    """A posting taken down during the interviews stays in the tab's dropdown."""
    save_interview_set(USER, answer_set(job_id))
    context = f"/api/interview/context?user={USER}"

    live = client.get(context).json()
    Database(config.user_db_path(USER)).mark_expired(job_id, "posting taken down")
    later = client.get(context).json()
    theirs = client.get(f"/api/interview/context?user={OTHER}").json()

    assert [job["id"] for job in live["jobs"]] == [job_id]
    assert live["saved_jobs"] == []
    assert later["jobs"] == []
    assert later["saved_jobs"] == [
        {
            "id": job_id,
            "title": TITLE,
            "company": COMPANY,
            "status": "expired",
            "fit_score": 80,
        }
    ]
    assert theirs["saved_jobs"] == []


@pytest.mark.parametrize(
    ("offset", "detail"),
    [(1, "does not match"), (0, "no longer exists")],
)
def test_edits_for_the_wrong_or_a_vanished_vacancy_are_refused(
    job_id: int, client: TestClient, offset: int, detail: str
) -> None:
    """An edit is saved only under the vacancy it belongs to, if that still exists."""
    target = job_id + 99
    body = answer_set(target).model_dump(mode="json")

    response = client.put(
        f"/api/interview/saved/answers/{target + offset}?user={USER}", json=body
    )

    assert response.status_code == 400
    assert detail in response.json()["detail"]


def test_the_questions_download_as_a_word_file(job_id: int, client: TestClient) -> None:
    """The body is what is on screen; the file opens and holds every question."""
    body = question_set(job_id).model_dump(mode="json")

    response = client.post(
        f"/api/interview/export/questions?user={USER}&format=docx", json=body
    )

    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == DOCX_TYPE
    assert response.headers["content-disposition"] == (
        f'attachment; filename="20260918 Interviewvragen {COMPANY}.docx"'
    )
    assert response.headers["cache-control"] == "no-store"
    text = [p.text for p in Document(io.BytesIO(response.content)).paragraphs]
    assert text[0] == f"{TITLE} bij {COMPANY}"
    assert QUESTION in text
    assert "Gebaseerd op: company review: cons" in text


def test_the_answers_download_with_the_applicants_own_edits(
    job_id: int, client: TestClient
) -> None:
    """A download takes the rewritten answer, not the draft that was generated."""
    body = answer_set(job_id, EDITED).model_dump(mode="json")

    word = client.post(f"/api/interview/export/answers?user={USER}", json=body)
    text = client.post(
        f"/api/interview/export/answers?user={USER}&format=txt", json=body
    )

    assert word.status_code == 200, word.text
    paragraphs = [p.text for p in Document(io.BytesIO(word.content)).paragraphs]
    assert EDITED in paragraphs and DRAFT not in paragraphs
    assert text.headers["content-type"] == "text/plain; charset=utf-8"
    assert text.headers["content-disposition"] == (
        f'attachment; filename="20260918 Interviewantwoorden {COMPANY}.txt"'
    )
    word_lines = [p.casefold() for p in paragraphs if p]
    text_lines = [line.casefold() for line in text.text.splitlines() if line]
    assert text_lines == word_lines


def test_a_company_name_cannot_break_the_download_header(
    job_id: int, client: TestClient
) -> None:
    """The name comes from a scraped page; the header must stay one safe line."""
    body = question_set(job_id).model_dump(mode="json")
    body["company"] = 'Evil"; filename="x.exe\r\nSet-Cookie: a=b'

    response = client.post(
        f"/api/interview/export/questions?user={USER}&format=txt", json=body
    )

    disposition = response.headers["content-disposition"]
    assert response.status_code == 200
    assert re.fullmatch(
        r'attachment; filename="[A-Za-z0-9 .,&()+-]+\.txt"', disposition
    )
    assert "set-cookie" not in response.headers


def test_a_control_character_cannot_cost_the_word_download(
    job_id: int, client: TestClient
) -> None:
    """Pasted or scraped control characters are dropped from both files alike."""
    pasted = Database(config.user_db_path(USER)).save_job(
        JobListing(
            title="Meet\x1bspecialist",
            company=COMPANY,
            url="https://voorbeeld.example/vacatures/8",
            source="board",
            description="Je bewaakt de kwaliteit van onze meetmethoden.",
            status=JobStatus.MATCHED,
            seen_at=datetime(2026, 1, 1, tzinfo=UTC),
        )
    )
    body = answer_set(pasted, "Eerste regel\x0bgeplakt uit Word\x00.").model_dump(
        mode="json"
    )
    body["company"] = "Find\x07where"
    export = f"/api/interview/export/answers?user={USER}"

    word = client.post(f"{export}&format=docx", json=body)
    text = client.post(f"{export}&format=txt", json=body)

    assert word.status_code == 200, word.text
    assert text.status_code == 200, text.text
    paragraphs = [p.text for p in Document(io.BytesIO(word.content)).paragraphs]
    assert paragraphs[0] == "Meetspecialist bij Findwhere"
    assert "Eerste regel\ngeplakt uit Word." in paragraphs
    assert "Eerste regel\ngeplakt uit Word." in text.text
    assert text.text.startswith("Meetspecialist bij Findwhere\n")
    assert not re.search("[\x00-\x08\x0b\x0c\x0e-\x1f]", text.text)


def test_an_unknown_format_or_an_empty_set_is_refused(
    job_id: int, client: TestClient
) -> None:
    """Only Word and text exist, and a set without questions is not a set."""
    body = question_set(job_id).model_dump(mode="json")
    empty = {**body, "questions": []}

    pdf = client.post(
        f"/api/interview/export/questions?user={USER}&format=pdf", json=body
    )
    nothing = client.post(f"/api/interview/export/questions?user={USER}", json=empty)

    assert pdf.status_code == 422
    assert nothing.status_code == 422


def test_saving_and_downloading_sit_behind_the_dashboard_token(
    job_id: int, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The routes live under /api/ precisely so the middleware covers them."""
    monkeypatch.setattr(
        "job_scout.web.app.load_secrets", lambda: {"dashboard_token": TOKEN}
    )
    body = question_set(job_id).model_dump(mode="json")
    saved = f"/api/interview/saved/{job_id}?user={USER}"
    export = f"/api/interview/export/questions?user={USER}"

    with TestClient(create_app()) as client:
        assert client.get(saved).status_code == 401
        assert client.post(export, json=body).status_code == 401

        client.headers["Authorization"] = f"Bearer {TOKEN}"
        assert client.get(saved).status_code == 200
        assert client.post(export, json=body).status_code == 200


def test_every_element_the_script_uses_is_on_the_page(client: TestClient) -> None:
    """A button the script wires up but the page lacks fails only in a browser."""
    script = client.get("/interview.js")
    page = (STATIC / "index.html").read_text(encoding="utf-8")

    assert script.status_code == 200
    ids = set(re.findall(r"el\('([a-z-]+)'\)", script.text))
    assert {"docx", "txt", "answers-docx", "answers-txt", "job", "language"} <= ids
    missing = [name for name in sorted(ids) if f'id="interview-{name}"' not in page]
    assert missing == []


def test_the_command_saves_what_it_generates(
    job_id: int, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A set generated on the command line shows up in the dashboard too."""
    monkeypatch.setattr("job_scout.cli.check_llm_available", lambda _: (True, None))
    monkeypatch.setattr("job_scout.cli.get_llm_client", lambda _: FakeLLMClient(["{}"]))
    monkeypatch.setattr(
        "job_scout.cli.generate_interview_answers", returning(answer_set(job_id))
    )

    result = CliRunner().invoke(
        cli, ["interview", "answers", str(job_id), "--user", USER]
    )
    saved = TestClient(create_app()).get(f"/api/interview/saved/{job_id}?user={USER}")

    assert result.exit_code == 0, result.output
    assert "Saved." in result.output
    assert saved.json()["answers"][0]["questions"][0]["draft_answer"] == DRAFT


@pytest.mark.parametrize(
    ("flags", "reply", "kept"),
    [
        ([], None, True),
        ([], "n\n", True),
        ([], "y\n", False),
        (["--yes"], None, False),
    ],
)
def test_the_command_asks_before_it_replaces_saved_answers(
    job_id: int,
    monkeypatch: pytest.MonkeyPatch,
    flags: list[str],
    reply: str | None,
    kept: bool,
) -> None:
    """Answers rewritten in the dashboard are not lost to a run on the command line.

    With no terminal to answer on, the saved answers are kept.
    """
    save_interview_set(USER, answer_set(job_id, EDITED))
    monkeypatch.setattr("job_scout.cli.check_llm_available", lambda _: (True, None))
    monkeypatch.setattr("job_scout.cli.get_llm_client", lambda _: FakeLLMClient(["{}"]))
    monkeypatch.setattr(
        "job_scout.cli.generate_interview_answers", returning(answer_set(job_id))
    )

    result = CliRunner().invoke(
        cli, ["interview", "answers", str(job_id), "--user", USER, *flags], input=reply
    )
    saved = TestClient(create_app()).get(f"/api/interview/saved/{job_id}?user={USER}")

    assert result.exit_code == 0, result.output
    draft = saved.json()["answers"][0]["questions"][0]["draft_answer"]
    assert draft == (EDITED if kept else DRAFT)
    assert ("Kept the saved answers." in result.output) is kept
    assert ("Replace them" in result.output) is not bool(flags)


def test_the_command_exports_a_saved_set_to_a_file(job_id: int, tmp_path: Path) -> None:
    """The export command writes the same Word file the dashboard offers."""
    save_interview_set(USER, answer_set(job_id, EDITED))
    folder = tmp_path / "out"
    folder.mkdir()

    result = CliRunner().invoke(
        cli,
        [
            "interview",
            "export",
            str(job_id),
            "--user",
            USER,
            "--mode",
            "answer",
            "--output",
            str(folder),
        ],
    )

    assert result.exit_code == 0, result.output
    written = folder / f"20260918 Interviewantwoorden {COMPANY}.docx"
    paragraphs = [p.text for p in Document(str(written)).paragraphs]
    assert EDITED in paragraphs


def test_the_command_says_when_nothing_is_saved(job_id: int, tmp_path: Path) -> None:
    """Exporting before generating is a clear error, not an empty file."""
    result = CliRunner().invoke(
        cli,
        [
            "interview",
            "export",
            str(job_id),
            "--user",
            USER,
            "--format",
            "txt",
            "--output",
            str(tmp_path / "vragen.txt"),
        ],
    )

    assert result.exit_code == 1
    assert "Nothing saved for vacancy" in result.output
    assert not (tmp_path / "vragen.txt").exists()
